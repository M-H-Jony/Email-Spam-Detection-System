import streamlit as st
import pickle
import string
import pandas as pd
import PyPDF2
import io
import docx
import openpyxl
from nltk.corpus import stopwords
import nltk
from nltk.stem.porter import Portermmer, PorterStemmer
from PyPDF2 import PdfReader

# Initialize NLTK resources
nltk.download('punkt')
nltk.download('stopwords')

ps = PorterStemmer()

# Function to extract text from PDF file
def extract_text_from_pdf(uploaded_file):
    pdf_reader = PdfReader(uploaded_file)
    text = ""
    for page_num in range(len(pdf_reader.pages)):
        text += pdf_reader.pages[page_num].extract_text()
    return text

# Function to extract text from DOCX file
def extract_text_from_docx(uploaded_file):
    doc = docx.Document(uploaded_file)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text
    return text

# Function to extract text from XLSX file
def extract_text_from_xlsx(uploaded_file):
    wb = openpyxl.load_workbook(uploaded_file)
    ws = wb.active
    text = ""
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                text += str(cell.value) + " "
    return text

# Function to transform text
def transform_text(text):
    text = text.lower()
    text = nltk.word_tokenize(text)

    y = []
    for i in text:
        if i.isalnum():
            y.append(i)

    text = y[:]
    y.clear()

    for i in text:
        if i not in stopwords.words('english') and i not in string.punctuation:
            y.append(i)

    text = y[:]
    y.clear()

    for i in text:
        y.append(ps.stem(i))

    return " ".join(y)

# Load TF-IDF vectorizer and model
tfidf = pickle.load(open('vectorizer.pkl', 'rb'))
model = pickle.load(open('model.pkl', 'rb'))

st.title("Email Spam Detection")

# Option for selecting input method (text or file upload)
option = st.radio("Select Input Method:", ("Text Input", "File Upload"))

if option == "Text Input":
    input_sms = st.text_area("Enter the mail content")

    if st.button('Check'):
        # Preprocess the text
        transformed_sms = transform_text(input_sms)

        # Vectorize the preprocessed text
        vector_input = tfidf.transform([transformed_sms])

        # Predict
        result = model.predict(vector_input)[0]

        # Display result
        if result == 1:
            st.header("Spam")
        else:
            st.header("Not Spam")

else:
    # File uploader for selecting input file
    uploaded_file = st.file_uploader("Choose a file", type=["txt", "pdf", "docx", "xlsx"])

    if uploaded_file is not None:
        if uploaded_file.type == "text/plain":
            # For TXT files
            file_contents = uploaded_file.getvalue().decode("utf-8")
        elif uploaded_file.type == "application/pdf":
            # For PDF files
            file_contents = extract_text_from_pdf(uploaded_file)
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            # For DOCX files
            file_contents = extract_text_from_docx(io.BytesIO(uploaded_file.read()))
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            # For XLSX files
            file_contents = extract_text_from_xlsx(uploaded_file)

        # Preprocess the text
        transformed_text = transform_text(file_contents)

        # Display the preprocessed text
        st.subheader("Preprocessed Text:")
        st.write(transformed_text)

        if st.button('Check'):
            # Vectorize the preprocessed text
            vector_input = tfidf.transform([transformed_text])

            # Predict
            result = model.predict(vector_input)[0]

            # Display result
            if result == 1:
                st.header("Spam")
            else:
                st.header("Not Spam")